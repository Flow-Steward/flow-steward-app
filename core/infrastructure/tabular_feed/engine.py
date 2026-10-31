"""Polars tabular feed engine ownership."""

from __future__ import annotations

import contextlib
import os
import tempfile
import time
from collections.abc import Mapping
from typing import Any

import psycopg

from core.application.extensions.artifact_dataset_plane_v2 import ArtifactDatasetPlane
from core.application.tabular_feed.contracts import snapshot_mode
from core.application.tabular_feed.snapshot_repository import (
    SnapshotConflictError,
    SnapshotRepository,
)
from core.infrastructure.date_types import DateLexicon, DateValueError
from core.infrastructure.extension_runtime.object_storage import APPLICATION_JSON_VALUE
from core.infrastructure.tabular_feed.actions import (
    _apply_rule as _actions_apply_rule,
)
from core.infrastructure.tabular_feed.actions import (
    _apply_rule_expressions,
    _normalize_actions,
    _validate_rule,
)
from core.infrastructure.tabular_feed.common import (
    _as_dict,
    _as_list,
    _envelope,
    _fail,
    _input,
    _limit_env,
    _neutralize_spreadsheet_formulas,
    _pl,
    _positive_int,
    _text,
)
from core.infrastructure.tabular_feed.conditions import _coerce_number, _empty
from core.infrastructure.tabular_feed.constants import (
    DEFAULT_MAX_MATERIALIZE_BYTES,
    DEFAULT_MAX_ROWS,
    DEFAULT_READ_ROWS_SCAN_MAX_BYTES,
    DEFAULT_READ_ROWS_SCAN_MAX_ROWS,
    PARQUET_CONTENT_TYPE,
    PREVIEW_LIMIT,
    READ_ROWS_CELL_BUDGET,
    READ_ROWS_LIMIT_MAX,
    STORAGE_CHUNK_SIZE,
    TEXT_CSV_VALUE,
)
from core.infrastructure.tabular_feed.dataframe_ops import _apply_column_map_df
from core.infrastructure.tabular_feed.engine_plane import (
    _artifact_meta_for_read,
    _candidate_meta_key,
    _create_artifact_from_file,
    _create_dataset,
    _create_dataset_from_df,
    _dataframe_from_artifact_path,
    _dataset_dataframe,
    _dataset_lazyframe,
    _dataset_meta_for_read,
    _df_rows,
    _df_schema,
    _handle_id,
    _meta_size_bytes,
    _plane,
    _preview,
    _read_dataset_rows,
    _read_json_key,
    _read_unfiltered_dataset_page,
    _require_read_grant,
    _require_scoped_owner,
    _schema,
    _schema_column_names,
    _write_json_key,
    injected_plane,
)
from core.infrastructure.tabular_feed.engine_snapshots import (
    _carry_forward_enrichment_lazy,
    _create_snapshot_parquet_artifact,
    _current_snapshot,
    _lazy_diff_projection,
    _lazy_diff_summary,
    _lazy_schema_names,
    _parse_snapshot_ref,
    _publication_snapshot_ref,
    _snapshot_read_metadata,
    _snapshot_repository,
    _snapshot_scope,
    _snapshot_scope_metadata,
    _validate_lazy_dataset_key,
    injected_snapshot_repository,
)
from core.infrastructure.tabular_feed.readers import (
    _format_from_meta,
)
from core.infrastructure.tabular_feed.relational import _apply_relational
from core.infrastructure.tabular_feed.storage_io import (
    _storage_key_path,
)

READ_FAILED = "tabular_feed.read_failed"
_FORMAT_LABELS = {
    "csv": "CSV",
    "json": "JSON",
    "ndjson": "NDJSON",
    "parquet": "Parquet",
    "tsv": "TSV",
    "xlsx": "XLSX",
}
_READ_FAILURE_DETAIL_CHARS = 300


class FeedReadError(Exception):
    """The input file could not be parsed in the format it was read as."""


@contextlib.contextmanager
def _reading_feed(fmt: str):
    """Report a parse failure of the input file as that, not as a raw Polars error.

    Polars raises ``ComputeError`` and its siblings for a malformed line or a value
    that does not fit the column type, sometimes only when a lazy scan is
    collected. Uncaught, it ended the run with the library's multi-line text and
    no step error, so the job page blamed the workflow's guardrails.
    """
    try:
        yield
    except _pl().exceptions.PolarsError as exc:
        label = _FORMAT_LABELS.get(fmt, fmt.upper() or "input")
        detail = next((line.strip() for line in str(exc).splitlines() if line.strip()), "")
        if len(detail) > _READ_FAILURE_DETAIL_CHARS:
            detail = detail[: _READ_FAILURE_DETAIL_CHARS - 3].rstrip() + "..."
        message = f"Could not read the {label} file"
        raise FeedReadError(f"{message}: {detail}" if detail else message) from exc


def _apply_rule(
    row: dict[str, Any],
    rule: dict[str, Any],
    *,
    date_lexicons: Mapping[str, DateLexicon] | None = None,
) -> bool:
    return _actions_apply_rule(row, rule, date_lexicons=date_lexicons)


def checksum(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    profile_key = _text(payload.get("profile_key"))
    raw_sha256 = _text(payload.get("raw_sha256"))
    if not profile_key or not raw_sha256:
        return _fail("invalid_payload", "profile_key and raw_sha256 are required")
    current = _current_snapshot(action, profile_key)
    metadata = current.metadata
    committed = _text(metadata.get("raw_sha256"))
    return _envelope(
        ok=True,
        result={
            "unchanged": bool(committed and committed == raw_sha256),
            "committed_snapshot_artifact_id": current.ref.snapshot_artifact_id,
            "committed_snapshot_ref": current.ref.as_dict(),
            "committed_raw_sha256": committed,
        },
    )


def inspect(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    options = _as_dict(payload.get("options"))
    meta, storage, max_bytes = _artifact_meta_for_read(
        action, _text(payload.get("artifact_handle"))
    )
    fmt = _format_from_meta(meta, options)
    suffix = ".jsonl" if fmt == "ndjson" else f".{fmt}"
    with _storage_key_path(
        storage, _text(meta.get("object_key")), suffix=suffix, max_bytes=max_bytes
    ) as path:
        with _reading_feed(fmt):
            df, info = _dataframe_from_artifact_path(path, meta, options)
            preview_rows = _df_rows(df, PREVIEW_LIMIT)
        columns = [
            {
                "name": column,
                "inferred_type": str(dtype),
                "sample_values": [
                    _text(row.get(column)) for row in preview_rows if not _empty(row.get(column))
                ][:3],
            }
            for column, dtype in zip(df.columns, df.dtypes, strict=True)
        ]
        return _envelope(
            ok=True,
            result={
                "format": info["format"],
                "encoding": info.get("encoding", ""),
                "delimiter": info.get("delimiter", ""),
                "sheets": info.get("sheets", []),
                "columns": columns,
                "row_count_estimate": df.height,
                "preview_rows": preview_rows,
                "warnings": info.get("warnings", []),
            },
        )


def read(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    options = _as_dict(payload.get("options"))
    meta, storage, max_bytes = _artifact_meta_for_read(
        action, _text(payload.get("artifact_handle"))
    )
    fmt = _format_from_meta(meta, options)
    suffix = ".jsonl" if fmt == "ndjson" else f".{fmt}"
    with (
        _storage_key_path(
            storage, _text(meta.get("object_key")), suffix=suffix, max_bytes=max_bytes
        ) as path,
        _reading_feed(fmt),
    ):
        df, info = _dataframe_from_artifact_path(path, meta, options, lazy=True)
        try:
            df = _apply_column_map_df(df, _as_dict(payload.get("column_map")))
        except ValueError as exc:
            return _fail("invalid_payload", str(exc), path="column_map")
        # A lazy scan reads the file here, so a malformed line surfaces here too.
        dataset_id, dataset_meta = _create_dataset_from_df(
            action,
            df=df,
            max_rows=info.get("max_rows"),
            source_metadata=_snapshot_read_metadata(
                action,
                meta,
                canonical=info["format"] == "parquet" and not _as_dict(payload.get("column_map")),
            ),
        )
        return _envelope(
            ok=True,
            result={
                "dataset_handle": dataset_id,
                "schema": dataset_meta["schema"],
                "row_count": dataset_meta["row_count"],
                "preview_rows": dataset_meta["preview_rows"],
                "warnings": info.get("warnings", []),
                "row_errors": {"count": 0, "dataset_handle": "", "sample": []},
            },
        )


def transform(
    action: dict[str, Any], *, date_lexicons: Mapping[str, DateLexicon] | None = None
) -> dict[str, Any]:
    payload = _input(action)
    with _dataset_lazyframe(action, _text(payload.get("dataset_handle"))) as (df, source_meta):
        rules = _normalize_actions(payload)
        try:
            normalized_rules = [_validate_rule(rule, index) for index, rule in enumerate(rules)]
        except ValueError as exc:
            return _fail("invalid_rules", str(exc), path="rules")
        try:
            out_df = _apply_column_map_df(df, _as_dict(payload.get("column_map")))
            out_df, report = _apply_rule_expressions(
                out_df,
                normalized_rules,
                row_count=int(source_meta.get("row_count") or 0),
                date_lexicons=date_lexicons,
            )
        except DateValueError as exc:
            path = "rules" if _as_list(payload.get("rules")) else "actions"
            return _fail(exc.code, str(exc), path=path)
        except (TypeError, ValueError) as exc:
            return _fail("invalid_rules", str(exc), path="rules")
        try:
            out_df = _apply_relational(out_df, _as_list(payload.get("relational")))
        except ValueError as exc:
            return _fail("invalid_rules", str(exc), path="relational")
        dataset_id, dataset_meta = _create_dataset_from_df(action, df=out_df)
    return _envelope(
        ok=True,
        result={
            "dataset_handle": dataset_id,
            "row_count": dataset_meta["row_count"],
            "preview_rows": dataset_meta["preview_rows"],
            "applied_rules_report": report,
            "warnings": [],
        },
    )


def concat(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    handles = [_text(item) for item in _as_list(payload.get("dataset_handles")) if _text(item)]
    if len(handles) < 2:
        return _fail("invalid_payload", "concat requires at least two dataset_handles")
    max_rows = _limit_env("FS_FEED_MAX_ROWS", DEFAULT_MAX_ROWS)
    max_bytes = _limit_env("FS_TABULAR_FEED_MAX_MATERIALIZE_BYTES", DEFAULT_MAX_MATERIALIZE_BYTES)
    metas = [_dataset_meta_for_read(action, handle) for handle in handles]
    total_rows = sum(
        _positive_int(meta.get("row_count")) or 0 for _dataset_id, meta, _storage in metas
    )
    total_bytes = sum(_meta_size_bytes(meta, storage) for _dataset_id, meta, storage in metas)
    if total_rows > max_rows:
        return _fail("too_many_rows", "concat input row_count exceeds operator maximum")
    if total_bytes > max_bytes:
        return _fail("feed_too_large", "concat input size exceeds operator maximum")
    column_map = _as_dict(payload.get("column_map"))
    try:
        frames = []
        for handle in handles:
            frame = _dataset_dataframe(action, handle, max_rows=max_rows, max_bytes=max_bytes)[0]
            applicable_map = {
                source: target
                for source, target in column_map.items()
                if _text(source) in frame.columns
            }
            frames.append(_apply_column_map_df(frame, applicable_map))
    except ValueError as exc:
        return _fail("invalid_payload", str(exc), path="column_map")
    first_schema = frames[0].schema
    for index, frame in enumerate(frames[1:], start=1):
        if frame.schema != first_schema:
            return _fail(
                "schema_mismatch",
                f"dataset_handles[{index}] schema does not match first dataset",
                path="dataset_handles",
            )
    df = _pl().concat(frames, how="vertical")
    try:
        df = _apply_relational(df, _as_list(payload.get("relational")))
    except ValueError as exc:
        return _fail("invalid_payload", str(exc), path="relational")
    dataset_id, dataset_meta = _create_dataset_from_df(action, df=df)
    return _envelope(
        ok=True,
        result={
            "dataset_handle": dataset_id,
            "row_count": df.height,
            "schema": dataset_meta.get("schema"),
            "preview_rows": _df_rows(df, PREVIEW_LIMIT),
        },
    )


def _type_valid(value: Any, expected_type: str) -> bool:
    if _empty(value):
        return True
    try:
        if expected_type in {"int", "integer"}:
            int(_coerce_number(value))
            return True
        if expected_type == "number":
            _coerce_number(value)
            return True
        if expected_type == "boolean":
            return _text(value).lower() in {"1", "0", "true", "false", "yes", "no", "on", "off"}
    except ValueError:
        return False
    return True


def validate(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    rows, meta = _read_dataset_rows(action, _text(payload.get("dataset_handle")))
    checks = _as_dict(payload.get("checks"))
    schema_names = _schema_column_names(_as_dict(meta.get("schema"))) or list(
        rows[0].keys() if rows else []
    )
    report: dict[str, Any] = {}

    required = [_text(item) for item in _as_list(checks.get("required_columns")) if _text(item)]
    missing = [column for column in required if column not in schema_names]
    report["required_columns"] = {
        "passed": not missing,
        "fail_count": len(missing),
        "sample_failures": missing[:PREVIEW_LIMIT],
    }

    non_empty = [_text(item) for item in _as_list(checks.get("non_empty")) if _text(item)]
    empty_failures = [
        {"row_index": index, "column": column}
        for index, row in enumerate(rows)
        for column in non_empty
        if _empty(row.get(column))
    ]
    report["non_empty"] = {
        "passed": not empty_failures,
        "fail_count": len(empty_failures),
        "sample_failures": empty_failures[:PREVIEW_LIMIT],
    }

    type_failures = []
    for column, expected in _as_dict(checks.get("types")).items():
        column_name = _text(column)
        expected_type = _text(expected)
        for index, row in enumerate(rows):
            if not _type_valid(row.get(column_name), expected_type):
                type_failures.append(
                    {"row_index": index, "column": column_name, "expected": expected_type}
                )
    report["types"] = {
        "passed": not type_failures,
        "fail_count": len(type_failures),
        "sample_failures": type_failures[:PREVIEW_LIMIT],
    }

    unique_failures = []
    for column in [_text(item) for item in _as_list(checks.get("unique")) if _text(item)]:
        seen: set[str] = set()
        for index, row in enumerate(rows):
            key = _text(row.get(column))
            if key in seen:
                unique_failures.append({"row_index": index, "column": column, "value": key})
            if key:
                seen.add(key)
    report["unique"] = {
        "passed": not unique_failures,
        "fail_count": len(unique_failures),
        "sample_failures": unique_failures[:PREVIEW_LIMIT],
    }
    failures = [
        failure
        for section in report.values()
        for failure in _as_list(section.get("sample_failures"))
    ]
    row_error_dataset = ""
    if len(failures) > 100:
        row_error_dataset, _ = _create_dataset(action, rows=failures)
    return _envelope(
        ok=True,
        result={
            "valid": all(section["passed"] for section in report.values()),
            "report": report,
            "row_errors": {
                "count": sum(section["fail_count"] for section in report.values()),
                "dataset_handle": row_error_dataset,
            },
        },
    )


def diff(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    handles = {}
    for field in ("dataset_handle", "baseline_dataset_handle"):
        value = payload.get(field)
        if not isinstance(value, str) or not value.strip():
            return _fail("invalid_payload", f"{field} is required", path=f"input.{field}")
        handles[field] = value.strip()
    profile_key = _text(payload.get("profile_key"))
    key_column = _text(payload.get("key_column"))
    if not profile_key or not key_column:
        return _fail("invalid_payload", "profile_key and key_column are required")
    try:
        purpose = snapshot_mode(
            payload.get("snapshot_mode"),
            has_reference=payload.get("expected_snapshot_ref") is not None,
        )
    except ValueError:
        return _fail("invalid_payload", "Invalid comparison purpose", path="snapshot_mode")
    expected_ref = (
        _parse_snapshot_ref(payload["expected_snapshot_ref"])
        if payload.get("expected_snapshot_ref") is not None
        else None
    )
    compare_columns = [
        _text(item) for item in _as_list(payload.get("compare_columns")) if _text(item)
    ]
    invalidate_on = [
        _text(item) for item in _as_list(payload.get("enrichment_invalidate_on")) if _text(item)
    ]
    with contextlib.ExitStack() as stack:
        datasets = {}
        for field, handle in handles.items():
            try:
                datasets[field] = stack.enter_context(_dataset_lazyframe(action, handle))
            except FileNotFoundError:
                return _fail("not_found", f"{field} dataset not found", path=f"input.{field}")
            except PermissionError:
                return _fail(
                    "access_denied", f"{field} dataset access denied", path=f"input.{field}"
                )
        current, dataset_meta = datasets["dataset_handle"]
        baseline, baseline_meta = datasets["baseline_dataset_handle"]
        if purpose == "publish":
            expected_ref = _publication_snapshot_ref(
                action,
                profile_key=profile_key,
                supplied=expected_ref,
                baseline=baseline,
                metadata=baseline_meta,
            )
        current = _validate_lazy_dataset_key(current, key_column=key_column)
        if not _lazy_schema_names(baseline):
            baseline = current.select(key_column).limit(0)
        baseline = _validate_lazy_dataset_key(baseline, key_column=key_column)
        current = _carry_forward_enrichment_lazy(
            current,
            baseline,
            key_column=key_column,
            invalidate_on=invalidate_on,
        )
        resolved_compare_columns = compare_columns or sorted(
            {
                column
                for column in [*_lazy_schema_names(current), *_lazy_schema_names(baseline)]
                if column != key_column and not column.startswith("__")
            }
        )
        projected = _lazy_diff_projection(
            current,
            baseline,
            key_column=key_column,
            compare_columns=resolved_compare_columns,
        )
        summary = _lazy_diff_summary(projected)

        diff_dataset_id, diff_meta = _create_dataset_from_df(
            action,
            df=projected.filter(_pl().col("__status") != "unchanged"),
        )
        snapshot_artifact_id, snapshot_meta = _create_snapshot_parquet_artifact(
            action,
            profile_key=profile_key,
            frame=current,
        )
        snapshot_schema = _df_schema(current)

    raw_sha256 = _text(payload.get("raw_sha256")) or _text(dataset_meta.get("sha256"))
    snapshot_metadata = {
        "profile_key": profile_key,
        "raw_sha256": raw_sha256,
        "row_count": int(dataset_meta.get("row_count") or 0),
        "key_column": key_column,
        "created_at": time.time(),
        "schema": snapshot_schema,
        "snapshot_artifact_id": snapshot_artifact_id,
        "artifact_object_key": snapshot_meta.get("object_key"),
        "snapshot_scope": _snapshot_scope_metadata(action, profile_key),
        "expected_snapshot_ref": expected_ref.as_dict() if expected_ref is not None else None,
    }
    _plane_obj, storage = _plane()
    _write_json_key(
        storage, _candidate_meta_key(action, profile_key, snapshot_artifact_id), snapshot_metadata
    )
    return _envelope(
        ok=True,
        result={
            "diff_dataset_handle": diff_dataset_id,
            "summary": summary,
            "snapshot_artifact_id": snapshot_artifact_id,
            "first_run": False,
            "preview_rows": diff_meta.get("preview_rows") or [],
        },
    )


def commit_snapshot(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    profile_key = _text(payload.get("profile_key"))
    snapshot_artifact_id = _handle_id(payload.get("snapshot_artifact_id"))
    if not profile_key or not snapshot_artifact_id:
        return _fail("invalid_payload", "profile_key and snapshot_artifact_id are required")
    _plane_obj, storage = _plane()
    _require_read_grant(action, kind="artifact", handle_id=snapshot_artifact_id)
    artifact_meta = _plane_obj.get_artifact(snapshot_artifact_id)
    if not artifact_meta:
        return _fail("not_found", "candidate snapshot artifact not found")
    _require_scoped_owner(action, artifact_meta, kind="artifact")
    if storage.stat(_text(artifact_meta.get("object_key"))) is None:
        return _fail("not_found", "candidate snapshot body not found")
    metadata = _read_json_key(
        storage, _candidate_meta_key(action, profile_key, snapshot_artifact_id)
    )
    if not metadata:
        return _fail("not_found", "candidate snapshot metadata not found")
    if metadata.get("expected_snapshot_ref") is None:
        return _fail("snapshot_baseline_required", "Candidate has no verified snapshot baseline")
    expected_ref = _parse_snapshot_ref(metadata["expected_snapshot_ref"])
    wanted_scope = _snapshot_scope_metadata(action, profile_key)
    if (
        metadata.get("snapshot_scope") != wanted_scope
        or artifact_meta.get("_tabular_snapshot_scope") != wanted_scope
        or metadata.get("snapshot_artifact_id") != snapshot_artifact_id
        or metadata.get("artifact_object_key") != artifact_meta.get("object_key")
    ):
        return _fail("snapshot_baseline_mismatch", "Candidate identity or scope does not match")
    _current_snapshot(action, profile_key)
    receipt = _snapshot_repository().commit(
        _snapshot_scope(action, profile_key),
        candidate_artifact_id=snapshot_artifact_id,
        expected_ref=expected_ref,
        metadata=metadata,
    )
    return _envelope(
        ok=True,
        result={
            "committed": True,
            "snapshot_artifact_id": snapshot_artifact_id,
            "previous_snapshot_artifact_id": receipt.previous_snapshot_artifact_id,
            "committed_snapshot_ref": receipt.ref.as_dict(),
        },
    )


def _row_matches_filter(row: dict[str, Any], filters: dict[str, Any]) -> bool:
    statuses = {_text(item) for item in _as_list(filters.get("status_in")) if _text(item)}
    if statuses and _text(row.get("__status")) not in statuses:
        return False
    for column, expected in _as_dict(filters.get("column_equals")).items():
        if row.get(_text(column)) != expected and _text(row.get(_text(column))) != _text(expected):
            return False
    return all(_empty(row.get(_text(column))) for column in _as_list(filters.get("missing")))


def _filtered_dataset_page(
    action: dict[str, Any],
    dataset_handle: str,
    *,
    filters: dict[str, Any],
    columns: list[str],
    offset: int,
    limit: int,
    max_rows: int,
    max_bytes: int,
) -> tuple[list[dict[str, Any]], int]:
    pl = _pl()
    with _dataset_lazyframe(
        action,
        dataset_handle,
        max_rows=max_rows,
        max_bytes=max_bytes,
    ) as (frame, _meta):
        names = set(_lazy_schema_names(frame))
        predicates = []
        statuses = [_text(item) for item in _as_list(filters.get("status_in")) if _text(item)]
        if statuses:
            predicates.append(
                pl.col("__status").is_in(statuses) if "__status" in names else pl.lit(False)
            )
        for column, expected in _as_dict(filters.get("column_equals")).items():
            name = _text(column)
            predicates.append(
                pl.col(name).cast(pl.String, strict=False).eq(_text(expected))
                if name in names
                else pl.lit(False)
            )
        for column in [_text(item) for item in _as_list(filters.get("missing")) if _text(item)]:
            if column not in names:
                predicates.append(pl.lit(True))
                continue
            value = pl.col(column)
            predicates.append(
                value.is_null() | (value.cast(pl.String, strict=False).str.strip_chars() == "")
            )
        if predicates:
            frame = frame.filter(pl.all_horizontal(predicates))
        total_matching = int(frame.select(pl.len()).collect(engine="streaming").item())
        page = frame.slice(offset, limit)
        if columns:
            page = page.select(
                [
                    pl.col(column) if column in names else pl.lit(None).alias(column)
                    for column in columns
                ]
            )
        rows = _df_rows(page.collect(engine="streaming"))
    return rows, total_matching


def read_rows(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    filters = _as_dict(payload.get("filter"))
    columns = [_text(item) for item in _as_list(payload.get("columns")) if _text(item)]
    offset = max(0, int(payload.get("offset") or 0))
    limit = min(READ_ROWS_LIMIT_MAX, max(1, int(payload.get("limit") or 100)))
    max_rows = _limit_env(
        "FS_TABULAR_FEED_READ_ROWS_SCAN_MAX_ROWS", DEFAULT_READ_ROWS_SCAN_MAX_ROWS
    )
    max_bytes = _limit_env(
        "FS_TABULAR_FEED_READ_ROWS_SCAN_MAX_BYTES", DEFAULT_READ_ROWS_SCAN_MAX_BYTES
    )
    if not filters:
        page, _meta, total_matching = _read_unfiltered_dataset_page(
            action,
            _text(payload.get("dataset_handle")),
            offset=offset,
            limit=limit,
            max_rows=max_rows,
            max_bytes=max_bytes,
        )
        inferred_count = len(page[0]) if page else 1
        projected_count = len(columns) if columns else max(1, inferred_count)
        if limit * projected_count > READ_ROWS_CELL_BUDGET:
            return _fail(
                "payload_budget_exceeded", "read_rows limit x columns exceeds payload budget"
            )
        if columns:
            page = [{column: row.get(column) for column in columns} for row in page]
        next_offset = offset + len(page) if offset + len(page) < total_matching else None
        return _envelope(
            ok=True,
            result={
                "rows": page,
                "returned": len(page),
                "total_matching": total_matching,
                "next_offset": next_offset,
            },
        )

    page, total_matching = _filtered_dataset_page(
        action,
        _text(payload.get("dataset_handle")),
        filters=filters,
        columns=columns,
        offset=offset,
        limit=limit,
        max_rows=max_rows,
        max_bytes=max_bytes,
    )
    if columns:
        projected_count = len(columns)
    elif page:
        projected_count = max(1, len(page[0]))
    else:
        projected_count = 1
    if limit * projected_count > READ_ROWS_CELL_BUDGET:
        return _fail("payload_budget_exceeded", "read_rows limit x columns exceeds payload budget")
    next_offset = offset + len(page) if offset + len(page) < total_matching else None
    return _envelope(
        ok=True,
        result={
            "rows": page,
            "returned": len(page),
            "total_matching": total_matching,
            "next_offset": next_offset,
        },
    )


def _bounded_payload_rows(rows: list[Any]) -> list[dict[str, Any]]:
    if len(rows) > READ_ROWS_LIMIT_MAX:
        raise OverflowError("payload_budget_exceeded")
    output = [dict(row) for row in rows if isinstance(row, dict)]
    projected_count = max((len(row) for row in output), default=1)
    if len(output) * projected_count > READ_ROWS_CELL_BUDGET:
        raise OverflowError("payload_budget_exceeded")
    return output


def create_dataset(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    rows = _bounded_payload_rows(_as_list(payload.get("rows")))
    append_to = _text(payload.get("append_to"))
    if append_to:
        existing, _meta = _read_dataset_rows(
            action,
            append_to,
            max_rows=_limit_env("FS_FEED_MAX_ROWS", DEFAULT_MAX_ROWS),
            max_bytes=_limit_env(
                "FS_TABULAR_FEED_MAX_MATERIALIZE_BYTES", DEFAULT_MAX_MATERIALIZE_BYTES
            ),
        )
        rows = [*existing, *rows]
    schema = payload.get("schema") if isinstance(payload.get("schema"), dict) else None
    dataset_id, dataset_meta = _create_dataset(action, rows=rows, schema=schema)
    return _envelope(
        ok=True,
        result={
            "dataset_handle": dataset_id,
            "row_count": len(rows),
            "schema": dataset_meta.get("schema") or _schema(rows),
            "preview_rows": _preview(rows),
        },
    )


def merge_rows(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    key_column = _text(payload.get("key_column"))
    if not key_column:
        return _fail("invalid_payload", "key_column is required", path="key_column")
    rows, _meta = _read_dataset_rows(action, _text(payload.get("dataset_handle")))
    updates = _bounded_payload_rows(_as_list(payload.get("rows")))
    allow_new = bool(payload.get("allow_new_columns", True))
    by_key = {_text(row.get(key_column)): dict(row) for row in rows if _text(row.get(key_column))}
    existing_columns = {column for row in rows for column in row}
    unmatched: list[str] = []
    merged = 0
    for update in updates:
        key = _text(update.get("key_column_value") or update.get(key_column))
        target = by_key.get(key)
        if target is None:
            unmatched.append(key)
            continue
        columns_to_set = _as_dict(update.get("columns_to_set")) or {
            column: value
            for column, value in update.items()
            if column not in {"key_column_value", key_column}
        }
        for column, value in columns_to_set.items():
            column_name = _text(column)
            if column_name not in existing_columns and (
                not allow_new or not column_name.startswith("enriched_")
            ):
                raise ValueError(f"new column '{column_name}' must be enriched_*")
            target[column_name] = value
            existing_columns.add(column_name)
        merged += 1
    output = [by_key.get(_text(row.get(key_column)), dict(row)) for row in rows]
    dataset_id, _dataset_meta = _create_dataset(action, rows=output)
    return _envelope(
        ok=True,
        result={
            "dataset_handle": dataset_id,
            "merged": merged,
            "unmatched_keys": {"count": len(unmatched), "sample": unmatched[:100]},
            "row_count": len(output),
            "preview_rows": _preview(output),
        },
    )


def write(action: dict[str, Any]) -> dict[str, Any]:
    payload = _input(action)
    fmt = _text(payload.get("format"), "csv").lower()
    if fmt == "jsonl":
        fmt = "ndjson"
    if fmt not in {"csv", "tsv", "json", "ndjson", "parquet", "xlsx"}:
        return _fail("unsupported_format", f"unsupported output format: {fmt}", path="format")
    dataset_handle = _text(payload.get("dataset_handle"))
    options = _as_dict(payload.get("options"))
    columns = [_text(item) for item in _as_list(options.get("columns")) if _text(item)]
    encoding = _text(options.get("encoding"), "utf-8")
    filename = _text(options.get("filename")) or f"feed-export.{fmt}"
    fd, output_path = tempfile.mkstemp(suffix=f".{fmt}")
    os.close(fd)
    primary_output_path = output_path
    extra_paths: list[str] = []
    row_count = 0
    try:
        if fmt in {"csv", "tsv", "ndjson", "parquet"}:
            with _dataset_lazyframe(action, dataset_handle) as (frame, meta):
                names = set(_lazy_schema_names(frame))
                if columns:
                    frame = frame.select([column for column in columns if column in names])
                row_count = int(meta.get("row_count") or 0)
                if row_count <= 0:
                    row_count = int(frame.select(_pl().len()).collect(engine="streaming").item())
                if fmt in {"csv", "tsv"}:
                    delimiter = "\t" if fmt == "tsv" else _text(options.get("delimiter"), ",")
                    if delimiter == "\\t":
                        delimiter = "\t"
                    if options.get("neutralize_formulas", True) is not False:
                        frame = _neutralize_spreadsheet_formulas(frame)
                    frame.sink_csv(output_path, separator=delimiter)
                    content_type = "text/tab-separated-values" if fmt == "tsv" else TEXT_CSV_VALUE
                elif fmt == "ndjson":
                    frame.sink_ndjson(output_path)
                    content_type = "application/x-ndjson"
                else:
                    frame.sink_parquet(output_path, compression="zstd")
                    content_type = PARQUET_CONTENT_TYPE
            if fmt in {"csv", "tsv"} and encoding.lower() not in {"utf8", "utf-8"}:
                fd_encoded, encoded_path = tempfile.mkstemp(suffix=f".{fmt}")
                os.close(fd_encoded)
                extra_paths.append(encoded_path)
                with (
                    open(output_path, encoding="utf-8") as source,
                    open(encoded_path, "wb") as target,
                ):
                    for chunk in iter(lambda: source.read(STORAGE_CHUNK_SIZE), ""):
                        target.write(chunk.encode(encoding, errors="replace"))
                output_path = encoded_path
        else:
            df, _meta = _dataset_dataframe(action, dataset_handle)
            if columns:
                df = df.select([column for column in columns if column in df.columns])
            if fmt == "xlsx" and options.get("neutralize_formulas", True) is not False:
                # openpyxl stores a string that starts with "=" as a formula.
                df = _neutralize_spreadsheet_formulas(df)
            row_count = df.height
        if fmt == "json":
            df.write_json(output_path)
            content_type = APPLICATION_JSON_VALUE
        elif fmt == "xlsx":
            import openpyxl

            if row_count > 1_048_576:
                return _fail(
                    "xlsx_row_limit_exceeded",
                    "XLSX export supports at most 1,048,576 rows; use CSV or Parquet",
                    path="format",
                )
            workbook = openpyxl.Workbook(write_only=True)
            sheet = workbook.create_sheet(_text(options.get("sheet"), "Sheet1"))
            sheet.append(df.columns)
            for row in df.iter_rows():
                sheet.append(list(row))
            if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
                del workbook["Sheet"]
            workbook.save(output_path)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        artifact_id, meta = _create_artifact_from_file(
            action,
            filename=filename,
            content_type=content_type,
            retention_class=_text(options.get("retention_class"), "run_window"),
            path=output_path,
        )
    finally:
        for path in {primary_output_path, output_path, *extra_paths}:
            with contextlib.suppress(FileNotFoundError):
                os.unlink(path)
    return _envelope(
        ok=True,
        result={
            "artifact_handle": artifact_id,
            "size_bytes": int(meta.get("size_bytes") or 0),
            "row_count": row_count,
            "sha256": _text(meta.get("sha256")),
        },
    )


def execute_operation(
    action_id: str,
    action: dict[str, Any],
    *,
    plane: ArtifactDatasetPlane,
    snapshot_repository: SnapshotRepository | None = None,
    date_lexicons: Mapping[str, DateLexicon] | None = None,
) -> dict[str, Any]:
    handler = ACTION_HANDLERS.get(_text(action_id))
    if handler is None:
        return _fail("invalid_payload", f"unknown_action:{action_id}")
    payload = dict(action or {})
    payload["action_id"] = _text(action_id)
    with injected_plane(plane), injected_snapshot_repository(snapshot_repository):
        try:
            if _text(action_id) == "transform":
                return transform(payload, date_lexicons=date_lexicons)
            return handler(payload)
        except FeedReadError as exc:
            return _fail(READ_FAILED, str(exc))
        except SnapshotConflictError as exc:
            return _fail(exc.code, str(exc))
        except psycopg.Error:
            return _fail(
                "snapshot_storage_unavailable",
                "Snapshot publication outcome could not be confirmed; retry the same candidate",
            )
        except FileNotFoundError as exc:
            return _fail("not_found", str(exc))
        except OverflowError as exc:
            return _fail(str(exc), str(exc))
        except ValueError as exc:
            code = str(exc)
            if code in {
                "feed_too_large",
                "invalid_snapshot_ref",
                "snapshot_baseline_mismatch",
                "snapshot_baseline_required",
            }:
                return _fail(code, code)
            return _fail("invalid_payload", str(exc))
        except PermissionError as exc:
            return _fail("access_denied", str(exc))
        except RuntimeError as exc:
            code = str(exc)
            if code in {"snapshot_storage_unavailable", "snapshot_dev_data_transition_required"}:
                return _fail(code, code)
            if code in {"row_error_ratio_exceeded"} or code.startswith("duplicate_key:"):
                return _fail(code.split(":", 1)[0], code)
            return _fail("runtime_failed", str(exc))


ACTION_HANDLERS = {
    "checksum": checksum,
    "inspect": inspect,
    "read": read,
    "transform": transform,
    "concat": concat,
    "diff": diff,
    "commit_snapshot": commit_snapshot,
    "read_rows": read_rows,
    "validate": validate,
    "create_dataset": create_dataset,
    "merge_rows": merge_rows,
    "write": write,
}
