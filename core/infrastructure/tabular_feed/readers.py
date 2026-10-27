"""Tabular artifact reader and format detection helpers."""

from __future__ import annotations

import contextlib
import csv
import io
import os
import tempfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

from core.infrastructure.tabular_feed.common import _clamped_limit, _pl, _text
from core.infrastructure.tabular_feed.conditions import _empty
from core.infrastructure.tabular_feed.constants import (
    DEFAULT_MAX_FEED_BYTES,
    DEFAULT_MAX_ROW_ERROR_RATIO,
    DEFAULT_MAX_ROWS,
)

DEFAULT_XLSX_BATCH_ROWS = 5_000


def _decode_feed(body: bytes, encoding: str = "") -> tuple[str, str]:
    if encoding:
        return body.decode(encoding, errors="replace"), encoding
    for candidate in ("utf-8-sig", "utf-8"):
        try:
            return body.decode(candidate), candidate
        except UnicodeDecodeError:
            continue
    return body.decode("latin-1", errors="replace"), "latin-1"


def _sniff_delimiter(text: str, filename: str, hint: str = "") -> str:
    if hint:
        return "\t" if hint == "\\t" else hint
    if filename.lower().endswith(".tsv"):
        return "\t"
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        return str(dialect.delimiter or ",")
    except csv.Error:
        return "\t" if "\t" in sample and sample.count("\t") > sample.count(",") else ","


def _dedupe_headers(headers: list[str]) -> tuple[list[str], list[str]]:
    warnings: list[str] = []
    counts: dict[str, int] = {}
    result: list[str] = []
    for index, header in enumerate(headers):
        name = _text(header).lstrip("\ufeff") or f"column_{index + 1}"
        if name != _text(header):
            warnings.append(f"empty_header:{index + 1}")
        count = counts.get(name, 0) + 1
        counts[name] = count
        result.append(name if count == 1 else f"{name}_{count}")
        if count > 1:
            warnings.append(f"duplicate_header:{name}")
    return result, sorted(set(warnings))


def _parse_rows(body: bytes, meta: dict[str, Any], options: dict[str, Any]) -> dict[str, Any]:
    max_bytes = _clamped_limit(
        options, "max_feed_bytes", "FS_FEED_MAX_BYTES", DEFAULT_MAX_FEED_BYTES
    )
    if len(body) > max_bytes:
        raise ValueError("feed_too_large")
    text, encoding = _decode_feed(body, _text(options.get("encoding")))
    delimiter = _sniff_delimiter(text, _text(meta.get("filename")), _text(options.get("delimiter")))
    header_row = max(1, int(options.get("header_row") or 1))
    skip_rows = max(0, int(options.get("skip_rows") or 0))
    lines = io.StringIO(text)
    for _ in range(skip_rows + header_row - 1):
        next(lines, None)
    reader = csv.reader(lines, delimiter=delimiter)
    raw_headers = next(reader, [])
    headers, warnings = _dedupe_headers(raw_headers)
    rows: list[dict[str, Any]] = []
    row_errors: list[dict[str, Any]] = []
    max_rows = _clamped_limit(options, "max_rows", "FS_FEED_MAX_ROWS", DEFAULT_MAX_ROWS)
    max_ratio = float(options.get("max_row_error_ratio") or DEFAULT_MAX_ROW_ERROR_RATIO)
    for raw_index, values in enumerate(reader, start=header_row + skip_rows + 1):
        if len(rows) >= max_rows:
            raise OverflowError("too_many_rows")
        if len(values) != len(headers):
            row_errors.append(
                {
                    "row_number": raw_index,
                    "reason": "ragged_row",
                    "expected_columns": len(headers),
                    "actual_columns": len(values),
                    "row": values[:20],
                }
            )
            total_seen = len(rows) + len(row_errors)
            if total_seen and len(row_errors) / total_seen > max_ratio:
                raise RuntimeError("row_error_ratio_exceeded")
            continue
        rows.append(dict(zip(headers, values, strict=True)))
    fmt = "tsv" if delimiter == "\t" else "csv"
    return {
        "format": fmt,
        "encoding": encoding,
        "delimiter": delimiter,
        "columns": headers,
        "rows": rows,
        "row_errors": row_errors,
        "warnings": warnings,
    }


def _format_from_meta(meta: dict[str, Any], options: dict[str, Any]) -> str:
    explicit = _text(options.get("format")).lower()
    if explicit:
        return "tsv" if explicit == "tab" else explicit
    filename = _text(meta.get("filename")).lower()
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix in {"csv", "tsv", "xlsx", "json", "ndjson", "jsonl", "parquet"}:
        return "ndjson" if suffix == "jsonl" else suffix
    content_type = _text(meta.get("content_type")).lower()
    if "parquet" in content_type:
        return "parquet"
    if "json" in content_type:
        return "json"
    return "tsv" if filename.endswith(".tsv") else "csv"


def _normalize_headers(columns: list[str]) -> tuple[list[str], list[str]]:
    return _dedupe_headers([_text(column) for column in columns])


def _sample_text_from_path(path: str, encoding: str = "") -> tuple[str, str]:
    with open(path, "rb") as handle:
        return _decode_feed(handle.read(8192), encoding)


def _dataframe_from_artifact_path(
    path: str,
    meta: dict[str, Any],
    options: dict[str, Any],
    *,
    polars_factory: Callable[[], Any] | None = None,
    lazy: bool = False,
) -> tuple[Any, dict[str, Any]]:
    pl = polars_factory() if polars_factory is not None else _pl()
    max_bytes = _clamped_limit(
        options, "max_feed_bytes", "FS_FEED_MAX_BYTES", DEFAULT_MAX_FEED_BYTES
    )
    if Path(path).stat().st_size > max_bytes:
        raise ValueError("feed_too_large")
    fmt = _format_from_meta(meta, options)
    max_rows = _clamped_limit(options, "max_rows", "FS_FEED_MAX_ROWS", DEFAULT_MAX_ROWS)

    def _collect_bounded(lazy_frame: Any) -> Any:
        df = lazy_frame.limit(max_rows + 1).collect(engine="streaming")
        if df.height > max_rows:
            raise OverflowError("too_many_rows")
        return df

    warnings: list[str] = []
    lazy_result = False
    if fmt in {"csv", "tsv"}:
        text, encoding = _sample_text_from_path(path, _text(options.get("encoding")))
        delimiter = (
            "\t"
            if fmt == "tsv"
            else _sniff_delimiter(
                text, _text(meta.get("filename")), _text(options.get("delimiter"))
            )
        )
        lazy_frame = pl.scan_csv(
            path,
            separator=delimiter,
            infer_schema=False,
            ignore_errors=False,
            truncate_ragged_lines=False,
            encoding="utf8-lossy" if encoding.lower() not in {"utf-8", "utf8"} else "utf8",
        )
        if lazy:
            source_headers = list(lazy_frame.collect_schema().names())
            headers, warnings = _normalize_headers(source_headers)
            rename = {
                source: target
                for source, target in zip(source_headers, headers, strict=True)
                if source != target
            }
            df = lazy_frame.rename(rename) if rename else lazy_frame
            lazy_result = True
        else:
            df = _collect_bounded(lazy_frame)
            headers, warnings = _normalize_headers(df.columns)
            df.columns = headers
        info = {
            "format": fmt,
            "encoding": encoding,
            "delimiter": delimiter,
            "warnings": warnings,
        }
    elif fmt == "parquet":
        lazy_frame = pl.scan_parquet(path)
        df = lazy_frame if lazy else _collect_bounded(lazy_frame)
        lazy_result = lazy
        info = {"format": "parquet", "encoding": "", "delimiter": "", "warnings": []}
    elif fmt == "ndjson":
        # Types come from the whole file: a column null in the first rows (Polars samples
        # 100) and filled later, such as Shopify's template_suffix, is not typed Null.
        lazy_frame = pl.scan_ndjson(path, infer_schema_length=None)
        df = lazy_frame if lazy else _collect_bounded(lazy_frame)
        lazy_result = lazy
        info = {"format": "ndjson", "encoding": "utf-8", "delimiter": "", "warnings": []}
    elif fmt == "json":
        df = pl.read_json(path, infer_schema_length=None)
        info = {"format": "json", "encoding": "utf-8", "delimiter": "", "warnings": []}
    elif fmt == "xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            sheet_name = _text(options.get("sheet")) or _text(options.get("sheet_name"))
            sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
            header_row = max(1, int(options.get("header_row") or 1))
            batch_rows = _xlsx_batch_rows(options)
            rows_iter = sheet.iter_rows(values_only=True)
            for _ in range(header_row - 1):
                next(rows_iter, None)
            raw_headers = list(next(rows_iter, []) or [])
            headers, warnings = _dedupe_headers([_text(item) for item in raw_headers])
            parquet_paths: list[str] = []
            batch: list[dict[str, Any]] = []
            row_count = 0
            for values in rows_iter:
                row = {
                    headers[index]: "" if value is None else value
                    for index, value in enumerate(values[: len(headers)])
                }
                if not any(not _empty(value) for value in row.values()):
                    continue
                if row_count >= max_rows:
                    raise OverflowError("too_many_rows")
                batch.append(row)
                row_count += 1
                if len(batch) >= batch_rows:
                    _write_xlsx_parquet_batch(pl, batch, parquet_paths)
                    batch = []
            if batch:
                _write_xlsx_parquet_batch(pl, batch, parquet_paths)
            try:
                if parquet_paths:
                    lazy_frames = [pl.scan_parquet(path) for path in parquet_paths]
                    lazy_frame = lazy_frames[0] if len(lazy_frames) == 1 else pl.concat(lazy_frames)
                    df = _collect_bounded(lazy_frame)
                else:
                    df = pl.DataFrame({header: [] for header in headers})
                info = {
                    "format": "xlsx",
                    "encoding": "",
                    "delimiter": "",
                    "sheets": workbook.sheetnames,
                    "warnings": warnings,
                }
            finally:
                for parquet_path in parquet_paths:
                    with contextlib.suppress(FileNotFoundError):
                        Path(parquet_path).unlink()
        finally:
            workbook.close()
    else:
        raise ValueError(f"unsupported_format:{fmt}")
    if lazy_result:
        info["max_rows"] = max_rows
    elif df.height > max_rows:
        raise OverflowError("too_many_rows")
    return df, info


def _xlsx_batch_rows(options: dict[str, Any]) -> int:
    raw = _text(options.get("xlsx_batch_rows"))
    if not raw:
        return DEFAULT_XLSX_BATCH_ROWS
    try:
        value = int(raw)
    except ValueError:
        value = DEFAULT_XLSX_BATCH_ROWS
    return max(1, value)


def _write_xlsx_parquet_batch(
    pl: Any,
    rows: list[dict[str, Any]],
    parquet_paths: list[str],
) -> None:
    fd, parquet_path = tempfile.mkstemp(suffix=".parquet")
    os.close(fd)
    try:
        pl.DataFrame(rows, infer_schema_length=None).write_parquet(parquet_path)
    except Exception:
        with contextlib.suppress(FileNotFoundError):
            Path(parquet_path).unlink()
        raise
    parquet_paths.append(parquet_path)
